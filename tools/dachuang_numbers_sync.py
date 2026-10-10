# -*- coding: utf-8 -*-
"""大创申报材料数字同步（2026-10-10）——申报书/PPT 的工程体量数字对齐当前代码库实测。

实测口径（与 2026-10-05 对齐会话一致）：
- 后端行数 = m2_server/**/*.py 排除 tests/（含空行）        -> 39,627
- 前端行数 = web/src/**/*.{ts,tsx}（含空行）                -> 41,867
- 测试用例 = pytest --collect-only(2,652) + vitest list(461) -> 3,113
- 测试文件 = m2_server/tests(117) + web 前端(46)            -> 163
- 插件     = m2_server/plugins 目录                          -> 28
- 路由     = APIRouter 定义文件(39) / @router 装饰器(226)
- 评测引擎 tools/voicebank_eval.py 1,573 行 + 102 用例
- 声音保存 voicepreserve 1,215 + api 404 行 / 115 用例
- 盲听     blindtest 464 + api 187 行 / 36 用例；前端 BlindTestPage 457 + useBlindTest 289 = 746
- 水印     audio_watermark 552 + api 139 行 / 31 用例
- 录课     course_api 554 + course_video 300 行 / 52 用例
"""
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DOCX = ROOT / "docs" / "大创申报-2026-10" / "大创项目申报书.docx"
PPTX = ROOT / "docs" / "大创申报-2026-10" / "大创申报PPT.pptx"

# ---------- docx ----------
# 段落（前缀匹配 → 新文本）
DOCX_PARAS = [
    ("1. 前期基础扎实，可行性风险低。",
     "1. **前期基础扎实，可行性风险低。** 已独立完成可运行的桌面端实时语音系统（后端约 4.0 万行 Python、"
     "前端约 4.2 万行 TS/TSX、3,113 个自动化测试、28 个能力插件，完整 CI 门禁，已开源，详见 §4.1）；"
     "克隆引擎、客观指标计算、盲听平台均已就绪——立项后直接进入数据采集。"),
    ("3. 评测引擎、听测平台与实验协议已就绪。",
     "3. **评测引擎、听测平台与实验协议已就绪。** 评测引擎 1,573 行 + 102 个测试，评测协议已撰写完成，"
     "不依赖 GPU 即可复现；听测平台支持手机扫码盲听，被试无需到场。"),
    ("水印 / 盲听 / 随访 / 录课以单元测试与链路自检验证",
     "水印 / 盲听 / 随访 / 录课以单元测试与链路自检验证（用例数见上表），同属链路验证。"),
]

# 表格 cell：(表号, 行, 列) → 新文本（T9=系统平台，T10=既有能力；行 0=表头，T10 已删随访行）
DOCX_CELLS = {
    (9, 1, 1): "39,627 行（不含测试），39 个路由模块，226 个路由端点",
    (9, 2, 1): "41,867 行",
    (9, 3, 1): "28 个（含声音保存、溯源水印、盲听评测、讲稿录课）",
    (9, 4, 1): "3,113 个用例（后端 2,652 + 前端 461），163 个测试文件",
    (10, 1, 2): "1,215 + 404 行 / 115 个测试",
    (10, 2, 2): "1,573 行 + 102 个测试",
    (10, 4, 2): "464 + 187 + 746 行 / 36 个测试",
    (10, 5, 2): "552 + 139 行 / 31 个测试",
    (10, 6, 2): "554 + 300 行 / 52 个测试",
}

# ---------- pptx ----------
PPTX_PARAS = [
    ("① 前期基础扎实：后端 3.8 万行 / 前端 4 万行 / 27 插件 / 2,847 测试 / 已开源 → 可行性风险低",
     "① 前期基础扎实：后端 4.0 万行 / 前端 4.2 万行 / 28 插件 / 3,113 测试 / 已开源 → 可行性风险低"),
    ("③ 评测引擎与协议已就绪：1559 行 + 99 测试，立项后可直接进入数据采集",
     "③ 评测引擎与协议已就绪：1,573 行 + 102 测试，立项后可直接进入数据采集"),
    ("后端 37,872 行 · 前端 40,169 行 · 27 插件 · 2,847 测试",
     "后端 39,627 行 · 前端 41,867 行 · 28 插件 · 3,113 测试"),
    ("工程体量：后端 37,872 行 · 前端 40,169 行 · 自动化测试 2,847 个（149 文件） · 能力插件 27 个 · CI 门禁与发布流程齐全 · 已在 GitHub 开源",
     "工程体量：后端 39,627 行 · 前端 41,867 行 · 自动化测试 3,113 个（163 文件） · 能力插件 28 个 · CI 门禁与发布流程齐全 · 已在 GitHub 开源"),
]


def set_docx_para(p, text: str) -> None:
    from docx.oxml.ns import qn
    proto = next((r for r in p.runs if r.text.strip()), None)
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for i, seg in enumerate(text.split("**")):
        if not seg:
            continue
        run = p.add_run(seg)
        if i % 2 == 1:
            run.bold = True
        if proto is not None:
            run.font.name = proto.font.name
            run.font.size = proto.font.size
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = rPr.makeelement(qn("w:rFonts"), {})
                rPr.append(rFonts)
            src = proto._element.find(qn("w:rPr"))
            if src is not None:
                src_fonts = src.find(qn("w:rFonts"))
                if src_fonts is not None:
                    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
                        val = src_fonts.get(qn(attr))
                        if val:
                            rFonts.set(qn(attr), val)


def patch_docx() -> list:
    from docx import Document
    doc = Document(str(DOCX))
    missed = []
    n = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        for pref, new in DOCX_PARAS:
            if t.startswith(pref):
                set_docx_para(p, new)
                n += 1
                break
    for (ti, r, c), new in DOCX_CELLS.items():
        try:
            cell = doc.tables[ti].rows[r].cells[c]
        except IndexError:
            missed.append(f"T{ti}R{r}C{c}")
            continue
        set_docx_para(cell.paragraphs[0], new)
        n += 1
    doc.save(str(DOCX))
    print(f"docx: 改 {n} 处，missed={missed or '无'}")
    return missed


def patch_pptx() -> list:
    from pptx import Presentation
    prs = Presentation(str(PPTX))
    norm = lambda s: re.sub(r"\s+", "", s)
    lookup = {norm(o): n for o, n in PPTX_PARAS}
    hit = set()
    for slide in prs.slides:
        for sh in slide.shapes:
            frames = []
            if sh.has_text_frame:
                frames.append(sh.text_frame)
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        frames.append(c.text_frame)
            for tf in frames:
                for para in tf.paragraphs:
                    full = "".join(r.text for r in para.runs)
                    new = lookup.get(norm(full))
                    if new is not None and new != full:
                        para.runs[0].text = new
                        for r in para.runs[1:]:
                            r.text = ""
                        hit.add(norm(new))
    prs.save(str(PPTX))
    missed = [o[:30] for o, n in PPTX_PARAS if norm(n) not in hit]
    print(f"pptx: 命中 {len(hit)}/{len(PPTX_PARAS)}，missed={missed or '无'}")
    return missed


if __name__ == "__main__":
    m1: list = []
    try:
        m1 = patch_docx()
    except ModuleNotFoundError:
        print("docx 跳过：当前解释器无 python-docx（用系统 python 跑）")
    m2 = patch_pptx()
    sys.exit(1 if (m1 or m2) else 0)
